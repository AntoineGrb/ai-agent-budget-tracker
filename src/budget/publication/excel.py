"""Classeur Excel du foyer : onglet Synthèse et un onglet par mois (spec 09).

`construire_classeur` est pure (aucun fichier) ; `enregistrer` écrit le classeur d'un bloc, par un
fichier temporaire renommé. Le classeur est une projection régénérable du ledger, comme l'annonce
le bandeau en A1 de chaque onglet.

Montants en `Decimal` dans les cellules, jamais en `float` ; textes forcés en chaîne, pour qu'un
libellé commençant par `=` ne devienne pas une formule.
"""

import os
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.chart import BarChart, PieChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import Series as SerieGraphique
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.worksheet import Worksheet

from budget.agents.synthese import CommentaireSynthese
from budget.agregats import Vue
from budget.categories import Categorie, libelle_de
from budget.errors import PublicationError
from budget.ledger.syntheses import SyntheseEnregistree
from budget.models import Compte
from budget.publication.contenu import ContenuPublication, MoisPublie
from budget.restitution.export_csv import COLONNES

SYNTHESE = "Synthèse"
LIGNE_TABLEAUX = 3  # premier tableau de l'onglet Synthèse, sous le bandeau
MOIS_GRAPHIQUE = 12
LIGNES_GRAPHIQUE_SYNTHESE = 19  # hauteur du graphique de l'onglet Synthèse, en lignes
COLONNES_RECAP = (
    "Catégorie",
    "Montant",
    "Budget",
    "Écart budget",
    "Moyenne 3 mois",
    "Écart moyenne",
)
COLONNE_PERSO = len(COLONNES_RECAP) + 2  # bloc perso en H, après une colonne vide
COLONNE_GRAPHIQUES = 2 * len(COLONNES_RECAP) + 3  # O, à droite des deux blocs
LIGNES_GRAPHIQUE = 15  # hauteur d'un graphique de répartition, en lignes

FORMAT_MONTANT = '#,##0.00 "€"'
FORMAT_ECART = '[Red]+#,##0.00 "€";[Color10]-#,##0.00 "€";0.00 "€"'  # dépassement en rouge
FORMAT_DATE = "dd/mm/yyyy"
GRAS = Font(bold=True)
SOUS_CATEGORIE = Font(italic=True, size=9)
RETRAIT = Alignment(indent=1)

type Valeur = str | Decimal | date | None


def bandeau(maintenant: datetime) -> str:
    return (
        f"Projection générée par budget le {maintenant:%d/%m/%Y à %H:%M} à partir du ledger. "
        "Toute modification manuelle sera écrasée à la prochaine publication."
    )


def a_regenerer(mois: str) -> str:
    return f"Commentaire à régénérer : budget synthese --mois {mois}"


def construire_classeur(contenu: ContenuPublication, maintenant: datetime) -> Workbook:
    """Synthèse, puis les mois du plus récent au plus ancien."""
    classeur = Workbook()
    synthese = classeur.active
    assert synthese is not None
    synthese.title = SYNTHESE
    _onglet_synthese(synthese, contenu.syntheses, maintenant)
    for mois in sorted(contenu.mois, key=lambda m: m.mois, reverse=True):
        _onglet_mois(classeur.create_sheet(mois.mois), mois, maintenant)
    classeur.properties.creator = "budget"
    classeur.properties.created = maintenant
    return classeur


def enregistrer(classeur: Workbook, chemin: Path) -> Path:
    """Écrit le classeur ; un fichier existant n'est remplacé qu'une fois le nouveau complet."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    temporaire = chemin.with_name(f".{chemin.stem}.tmp.xlsx")
    try:
        classeur.save(temporaire)
        os.replace(temporaire, chemin)
    except PermissionError:
        temporaire.unlink(missing_ok=True)
        raise PublicationError(
            f"{chemin} n'a pas pu être remplacé : le fermer dans Excel, puis relancer "
            "`budget publish`"
        ) from None
    return chemin


# Cellules -------------------------------------------------------------------------------------


def _ecrire(onglet: Worksheet, ligne: int, colonne: int, valeur: Valeur, fmt: str = "") -> Cell:
    cellule: Cell = onglet.cell(row=ligne, column=colonne)
    if isinstance(valeur, str):
        cellule.value = valeur
        cellule.data_type = "s"  # jamais de formule, même pour un libellé `=…`
    else:
        cellule.value = valeur
    if fmt:
        cellule.number_format = fmt
    return cellule


def _ligne(
    onglet: Worksheet,
    ligne: int,
    colonne: int,
    valeurs: Sequence[Valeur],
    formats: Sequence[str] = (),
    *,
    police: Font | None = None,
) -> None:
    for i, valeur in enumerate(valeurs):
        fmt = formats[i] if i < len(formats) else ""
        cellule = _ecrire(onglet, ligne, colonne + i, valeur, fmt)
        if police is not None:
            cellule.font = police


def _entete(onglet: Worksheet, maintenant: datetime) -> None:
    _ecrire(onglet, 1, 1, bandeau(maintenant)).font = Font(italic=True)


def _resume(commentaire: CommentaireSynthese | None, mois: str) -> str:
    return commentaire.resume if commentaire is not None else a_regenerer(mois)


def _nom_vue(compte: Compte) -> str:
    return "Perso" if compte is Compte.PERSO else "Joint"


# Synthèse -------------------------------------------------------------------------------------


def colonnes_synthese() -> list[str]:
    """Colonnes d'un tableau de vue : mois, total, une colonne par catégorie."""
    return ["Mois", "Total", *(c.libelle for c in Categorie)]


def _onglet_synthese(
    onglet: Worksheet, syntheses: Sequence[SyntheseEnregistree], maintenant: datetime
) -> None:
    """Tableau joint, tableau perso en dessous, graphique, puis commentaires des mois."""
    _entete(onglet, maintenant)
    ligne = LIGNE_TABLEAUX
    premieres: dict[Compte, int] = {}
    for compte in Compte:
        premieres[compte] = _tableau_vue(onglet, compte, syntheses, ligne)
        ligne = premieres[compte] + len(syntheses) + 1  # une ligne vide entre deux blocs
    onglet.freeze_panes = "B1"
    onglet.column_dimensions["A"].width = 12
    for colonne in range(2, len(colonnes_synthese()) + 1):
        onglet.column_dimensions[get_column_letter(colonne)].width = 13
    if syntheses:
        _graphique_synthese(onglet, len(syntheses), premieres, ligne)
        ligne += LIGNES_GRAPHIQUE_SYNTHESE + 1
    _ecrire(onglet, ligne, 1, "Commentaires").font = GRAS
    _ligne(onglet, ligne + 1, 1, ["Mois", "Commentaire"], police=GRAS)
    for i, s in enumerate(syntheses):
        _ligne(onglet, ligne + 2 + i, 1, [s.mois, _resume(s.commentaire, s.mois)])


def _tableau_vue(
    onglet: Worksheet, compte: Compte, syntheses: Sequence[SyntheseEnregistree], ligne: int
) -> int:
    """Titre, en-têtes, une ligne par mois ; renvoie la ligne du premier mois."""
    _ecrire(onglet, ligne, 1, f"Vue {compte.value}").font = GRAS
    _ligne(onglet, ligne + 1, 1, colonnes_synthese(), police=GRAS)
    formats = ["", *[FORMAT_MONTANT] * (1 + len(Categorie))]
    for i, s in enumerate(syntheses):
        vue = s.agregats.vue(compte)
        par_categorie = {c.categorie: c.montant for c in vue.categories}
        valeurs: list[Valeur] = [
            s.mois,
            vue.total,
            *(par_categorie.get(c, Decimal("0.00")) for c in Categorie),
        ]
        _ligne(onglet, ligne + 2 + i, 1, valeurs, formats)
        onglet.cell(row=ligne + 2 + i, column=2).font = GRAS
    return ligne + 2


def _graphique_synthese(
    onglet: Worksheet, nombre: int, premieres: dict[Compte, int], ligne: int
) -> None:
    """Totaux des deux vues sur les 12 derniers mois, en colonnes groupées."""
    decalage = max(0, nombre - MOIS_GRAPHIQUE)
    graphique = BarChart()
    graphique.type = "col"
    graphique.grouping = "clustered"  # jamais empilé : les deux vues ne s'additionnent pas
    graphique.title = f"Dépenses mensuelles, {MOIS_GRAPHIQUE} derniers mois"
    graphique.y_axis.title = "Dépenses (€)"
    graphique.x_axis.title = "Mois"
    graphique.y_axis.numFmt = FORMAT_MONTANT
    graphique.y_axis.delete = False
    graphique.x_axis.delete = False

    def fenetre(compte: Compte, colonne: int) -> Reference:
        debut = premieres[compte] + decalage
        return Reference(
            onglet, min_col=colonne, min_row=debut, max_row=premieres[compte] + nombre - 1
        )

    for compte in Compte:
        graphique.series.append(Series(fenetre(compte, 2), title=f"{_nom_vue(compte)} : total"))
    graphique.set_categories(fenetre(Compte.JOINT, 1))
    graphique.width, graphique.height = 24, 9
    onglet.add_chart(graphique, f"A{ligne}")


# Onglet mensuel -------------------------------------------------------------------------------


def _onglet_mois(onglet: Worksheet, publie: MoisPublie, maintenant: datetime) -> None:
    """Commentaire, récapitulatifs et graphiques en haut ; opérations filtrables en dessous.

    Le filtre masque des lignes entières : rien d'autre ne partage les lignes des opérations.
    """
    _entete(onglet, maintenant)
    commentaire = publie.synthese.commentaire
    lignes = [_resume(commentaire, publie.mois)]
    if commentaire is not None:
        lignes += [f"• {point}" for point in commentaire.points_cles]
    for i, texte in enumerate(lignes):
        _ecrire(onglet, 3 + i, 1, texte)
    onglet.cell(row=3, column=1).font = GRAS

    debut = 3 + len(lignes) + 1
    agregats = publie.synthese.agregats
    hauteur = 0
    for rang, (compte, colonne) in enumerate(((Compte.JOINT, 1), (Compte.PERSO, COLONNE_PERSO))):
        vue = agregats.vue(compte)
        hauteur_vue, categories = _recapitulatif(onglet, vue, debut, colonne)
        hauteur = max(hauteur, hauteur_vue)
        ancre = f"{get_column_letter(COLONNE_GRAPHIQUES + rang * 8)}{debut}"
        _secteurs(onglet, vue, categories, colonne, ancre)

    entete = debut + max(hauteur, LIGNES_GRAPHIQUE) + 1
    _ligne(onglet, entete, 1, COLONNES, police=GRAS)
    formats = [FORMAT_DATE, FORMAT_DATE, "", "", FORMAT_MONTANT]
    for i, o in enumerate(publie.operations):
        valeurs: list[Valeur] = [
            o.date_operation,
            o.date_debit,
            o.compte,
            o.libelle,
            o.montant,
            o.categorie,
            o.sous_categorie,
            o.comptee,
            o.regle,
        ]
        _ligne(onglet, entete + 1 + i, 1, valeurs, formats)
    derniere = entete + len(publie.operations)
    onglet.auto_filter.ref = f"A{entete}:{get_column_letter(len(COLONNES))}{derniere}"
    # A et H portent aussi les catégories des récapitulatifs
    for lettre, largeur in zip("ABCDEFGHI", (20, 14, 13, 45, 13, 14, 22, 20, 16), strict=True):
        onglet.column_dimensions[lettre].width = largeur
    for colonne in range(COLONNE_PERSO + 1, COLONNE_PERSO + len(COLONNES_RECAP)):
        lettre = get_column_letter(colonne)
        onglet.column_dimensions[lettre].width = max(
            onglet.column_dimensions[lettre].width or 0, 14
        )


def _recapitulatif(onglet: Worksheet, vue: Vue, ligne: int, colonne: int) -> tuple[int, list[int]]:
    """Bloc d'une vue : titre, en-têtes, chaque catégorie suivie de ses sous-catégories, total.

    Renvoie la hauteur du bloc et les lignes des catégories, sources des secteurs.
    """
    _ecrire(onglet, ligne, colonne, f"Vue {vue.compte.value}").font = GRAS
    _ligne(onglet, ligne + 1, colonne, COLONNES_RECAP, police=GRAS)
    formats = [
        "",
        FORMAT_MONTANT,
        FORMAT_MONTANT,
        FORMAT_ECART,
        FORMAT_MONTANT,
        FORMAT_ECART,
    ]
    courante = ligne + 2
    categories = []
    for c in vue.categories:
        valeurs: list[Valeur] = [
            c.categorie.libelle,
            c.montant,
            c.budget,
            c.ecart_budget,
            c.moyenne_3_mois,
            c.ecart_moyenne,
        ]
        _ligne(onglet, courante, colonne, valeurs, formats)
        categories.append(courante)
        courante += 1
        for s in c.sous_categories:  # non nulles, par montant décroissant
            detail: list[Valeur] = [libelle_de(s.sous_categorie), s.montant]
            _ligne(onglet, courante, colonne, detail, formats, police=SOUS_CATEGORIE)
            onglet.cell(row=courante, column=colonne).alignment = RETRAIT
            courante += 1
    _ligne(onglet, courante, colonne, ["Total", vue.total], formats, police=GRAS)
    return courante - ligne + 1, categories


def _secteurs(onglet: Worksheet, vue: Vue, lignes: Sequence[int], colonne: int, ancre: str) -> None:
    """Répartition par catégorie ; une catégorie de montant négatif en est exclue.

    Les lignes des catégories sont séparées par leurs sous-catégories : séries en référence
    d'union, `('2026-07'!$B$9,'2026-07'!$B$12)`.
    """
    retenues = [n for n, c in zip(lignes, vue.categories, strict=True) if c.montant >= 0]
    if not retenues:
        return
    graphique = PieChart()
    graphique.title = f"Vue {vue.compte.value} : répartition par catégorie"
    graphique.dataLabels = DataLabelList(showPercent=True)
    graphique.width, graphique.height = 13, 7.5
    serie = SerieGraphique()
    serie.val = NumDataSource(numRef=NumRef(f=_union(onglet, colonne + 1, retenues)))
    serie.cat = AxDataSource(strRef=StrRef(f=_union(onglet, colonne, retenues)))
    graphique.series.append(serie)
    onglet.add_chart(graphique, ancre)


def _union(onglet: Worksheet, colonne: int, lignes: Sequence[int]) -> str:
    feuille = quote_sheetname(onglet.title)
    lettre = get_column_letter(colonne)
    return "(" + ",".join(f"{feuille}!${lettre}${ligne}" for ligne in lignes) + ")"
