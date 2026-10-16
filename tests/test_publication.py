"""Classeur Excel de juin et juillet (spec 09) : onglets, valeurs, formats, graphiques, contrôles
bloquants et écriture du fichier."""

import os
from collections.abc import Iterator
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from budget import cli
from budget.config import Contexte, charger_contexte
from budget.errors import ControleError, PublicationError
from budget.importation import importer_mois
from budget.ledger import Ledger
from budget.publication.contenu import ContenuPublication
from budget.publication.excel import (
    FORMAT_DATE,
    FORMAT_ECART,
    FORMAT_MONTANT,
    LIGNES_GRAPHIQUE,
    SYNTHESE,
    a_regenerer,
    construire_classeur,
    enregistrer,
)
from budget.restitution.export_csv import COLONNES
from tests.conftest import DOSSIER_FIXTURES, EXEMPLE_CONTEXTE
from tests.scenarios.publication import MAINTENANT, contenu, preparer_ledger

BANDEAU = (
    "Projection générée par budget le 16/10/2026 à 18:40 à partir du ledger. Toute modification "
    "manuelle sera écrasée à la prochaine publication."
)
LIGNE_RECAP = 7  # juillet : commentaire en lignes 3 à 5, récapitulatifs à partir de la ligne 7


@pytest.fixture(scope="module")
def contexte() -> Contexte:
    return charger_contexte(EXEMPLE_CONTEXTE)


@pytest.fixture(scope="module")
def ledger(tmp_path_factory: pytest.TempPathFactory, contexte: Contexte) -> Iterator[Ledger]:
    chemin = tmp_path_factory.mktemp("publication") / "budget.sqlite"
    with preparer_ledger(chemin, contexte) as ledger:
        yield ledger


@pytest.fixture(scope="module")
def juillet(ledger: Ledger, contexte: Contexte) -> ContenuPublication:
    return contenu(ledger, contexte)


@pytest.fixture(scope="module")
def classeur(juillet: ContenuPublication) -> Workbook:
    return construire_classeur(juillet, MAINTENANT)


def _ligne(onglet: Any, numero: int, debut: int = 1, fin: int = 12) -> list[object]:
    return [onglet.cell(row=numero, column=c).value for c in range(debut, fin + 1)]


def _operations(onglet: Any) -> tuple[int, int]:
    """Ligne d'en-tête et dernière ligne des opérations, d'après le filtre."""
    debut, fin = onglet.auto_filter.ref.split(":")
    return int(debut[1:]), int(fin[1:])


def test_onglets_synthese_puis_mois_recents(classeur: Workbook) -> None:
    assert classeur.sheetnames == [SYNTHESE, "2026-07", "2026-06"]
    for onglet in classeur.worksheets:
        assert onglet["A1"].value == BANDEAU
    assert [len(o._charts) for o in classeur.worksheets] == [1, 2, 2]


def test_aucun_float_et_textes_jamais_formules(classeur: Workbook) -> None:
    for onglet in classeur.worksheets:
        for ligne in onglet.iter_rows():
            for cellule in ligne:
                assert not isinstance(cellule.value, float), (onglet.title, cellule.coordinate)
                if isinstance(cellule.value, str):
                    assert cellule.data_type == "s", (onglet.title, cellule.coordinate)


def test_libelle_commencant_par_egal_reste_du_texte(
    juillet: ContenuPublication, tmp_path: Path
) -> None:
    mois = juillet.mois[-1]
    operations = [replace(mois.operations[0], libelle="=SOMME(A1:A9)"), *mois.operations[1:]]
    piege = replace(juillet, mois=[*juillet.mois[:-1], replace(mois, operations=operations)])
    chemin = enregistrer(construire_classeur(piege, MAINTENANT), tmp_path / "budget.xlsx")
    relu = load_workbook(chemin)["2026-07"]
    entete, _ = _operations(relu)
    cellule = relu.cell(row=entete + 1, column=4)
    assert (cellule.value, cellule.data_type) == ("=SOMME(A1:A9)", "s")


def test_synthese_tableau_joint_puis_perso(classeur: Workbook) -> None:
    onglet = classeur[SYNTHESE]
    colonnes = ["Mois", "Total", "Logement", "Courses", "Restos", "Loisirs", "Achats"]
    colonnes += ["Enfants", "Transport", "Abonnements", "Santé", "Divers"]
    # Joint : titre ligne 3, en-têtes ligne 4, juin et juillet ; perso cinq lignes plus bas
    assert onglet["A3"].value == "Vue joint" and onglet["A8"].value == "Vue perso"
    for entetes in (4, 9):
        assert _ligne(onglet, entetes, fin=13) == [*colonnes, None]
    joint_juin, joint_juillet = _ligne(onglet, 5), _ligne(onglet, 6)
    perso_juin, perso_juillet = _ligne(onglet, 10), _ligne(onglet, 11)
    assert _ligne(onglet, 7) == [None] * 12
    assert [joint_juin[0], joint_juillet[0], perso_juin[0], perso_juillet[0]] == [
        "2026-06",
        "2026-07",
        "2026-06",
        "2026-07",
    ]
    colonne = {nom: i for i, nom in enumerate(colonnes)}
    assert [joint_juin[colonne[n]] for n in ("Total", "Logement", "Divers")] == [
        Decimal("7566.46"),
        Decimal("3289.26"),
        Decimal("0.00"),
    ]
    assert [perso_juin[colonne[n]] for n in ("Total", "Loisirs")] == [
        Decimal("1322.86"),
        Decimal("685.69"),
    ]
    assert [joint_juillet[colonne[n]] for n in ("Total", "Divers")] == [
        Decimal("9404.27"),
        Decimal("618.45"),
    ]
    assert perso_juillet[colonne["Total"]] == Decimal("1154.49")
    assert onglet["B5"].number_format == FORMAT_MONTANT and onglet["B5"].font.b
    assert onglet.freeze_panes == "B1"

    # Graphique sous les deux tableaux, séries des deux totaux, en colonnes groupées
    [graphique] = onglet._charts
    assert graphique.anchor == "A13"
    assert (graphique.type, graphique.grouping) == ("col", "clustered")
    assert [s.tx.v for s in graphique.series] == ["Joint : total", "Perso : total"]
    assert [s.val.numRef.f for s in graphique.series] == [
        "'Synthèse'!$B$5:$B$6",
        "'Synthèse'!$B$10:$B$11",
    ]

    # Commentaires sous le graphique
    assert onglet["A33"].value == "Commentaires"
    assert _ligne(onglet, 34, fin=2) == ["Mois", "Commentaire"]
    assert onglet["A36"].value == "2026-07"
    assert str(onglet["B36"].value).startswith("Joint : 9 404,27 €")


def test_synthese_sans_commentaire_a_regenerer(juillet: ContenuPublication) -> None:
    sans = replace(
        juillet,
        mois=[
            replace(m, synthese=m.synthese.model_copy(update={"commentaire": None}))
            if m.mois == "2026-06"
            else m
            for m in juillet.mois
        ],
    )
    classeur = construire_classeur(sans, MAINTENANT)
    assert classeur[SYNTHESE]["B35"].value == a_regenerer("2026-06")
    assert a_regenerer("2026-06") == "Commentaire à régénérer : budget synthese --mois 2026-06"
    assert classeur["2026-06"]["A3"].value == a_regenerer("2026-06")


def _bloc(onglet: Any, ligne: int, colonne: int) -> list[tuple[list[object], bool]]:
    """Lignes d'un récapitulatif, du titre au Total ; vrai pour une ligne de sous-catégorie."""
    lignes = []
    while True:
        valeurs = _ligne(onglet, ligne, colonne, colonne + 5)
        cellule = onglet.cell(row=ligne, column=colonne)
        lignes.append((valeurs, bool(cellule.font.i)))
        if valeurs[0] == "Total":
            return lignes
        ligne += 1


def _categories(
    bloc: list[tuple[list[object], bool]],
) -> dict[object, tuple[list[object], list[list[object]]]]:
    """Catégorie → (sa ligne, ses lignes de sous-catégories)."""
    resultat: dict[object, tuple[list[object], list[list[object]]]] = {}
    courante: list[list[object]] = []
    for valeurs, sous in bloc[2:-1]:
        if sous:
            courante.append(valeurs)
        else:
            courante = []
            resultat[valeurs[0]] = (valeurs, courante)
    return resultat


def test_onglet_mensuel_de_juillet(classeur: Workbook) -> None:
    onglet = classeur["2026-07"]
    assert [onglet.cell(row=n, column=1).value for n in (3, 4, 5, 6)] == [
        "Joint : 9 404,27 €, surtout Logement (4 416,17 €). Perso : 1 154,49 €.",
        "• Courses joint : +157,92 € de budget.",
        "• Divers joint : 618,45 €, nouveau ce mois.",
        None,
    ]
    joint, perso = _bloc(onglet, LIGNE_RECAP, 1), _bloc(onglet, LIGNE_RECAP, 8)
    assert joint[0][0][0] == "Vue joint" and perso[0][0][0] == "Vue perso"
    assert joint[1][0] == [
        "Catégorie",
        "Montant",
        "Budget",
        "Écart budget",
        "Moyenne 3 mois",
        "Écart moyenne",
    ]
    for bloc, total in ((joint, "9404.27"), (perso, "1154.49")):
        assert bloc[-1][0][:2] == ["Total", Decimal(total)]
        categories = _categories(bloc)
        assert sum(ligne[1] for ligne, _ in categories.values()) == Decimal(total)  # type: ignore[misc]
        for ligne, sous in categories.values():
            # Chaque catégorie suivie de ses sous-catégories, dont la somme fait la catégorie
            assert sous and sum(s[1] for s in sous) == ligne[1]  # type: ignore[misc]
            assert all(s[2:] == [None] * 4 for s in sous)

    par_categorie = _categories(joint)
    courses, _ = par_categorie["Courses"]
    assert courses == [
        "Courses",
        Decimal("757.92"),
        Decimal("600.00"),
        Decimal("157.92"),
        Decimal("627.36"),
        Decimal("130.56"),
    ]
    logement, detail = par_categorie["Logement"]
    assert logement[4] == Decimal("3289.26")
    assert ["Autre", Decimal("2659.16")] in [s[:2] for s in detail]  # acompte climatisation
    assert par_categorie["Santé"][0][2:4] == [None, None]  # ni budget ni écart

    ligne_courses = next(
        n for n in range(LIGNE_RECAP, 40) if onglet.cell(row=n, column=1).value == "Courses"
    )
    assert onglet.cell(row=ligne_courses, column=4).number_format == FORMAT_ECART
    assert onglet.cell(row=ligne_courses, column=2).number_format == FORMAT_MONTANT
    sous_categorie = onglet.cell(row=ligne_courses + 1, column=1)
    assert (sous_categorie.font.i, sous_categorie.font.sz, sous_categorie.alignment.indent) == (
        True,
        9,
        1,
    )
    assert onglet.cell(row=ligne_courses + 1, column=2).font.i

    entete, derniere = _operations(onglet)
    assert derniere - entete == 169
    assert _ligne(onglet, entete, 1, 9) == list(COLONNES)
    operations = [_ligne(onglet, n, 1, 9) for n in range(entete + 1, derniere + 2)]
    assert all(o[0] for o in operations[:169]) and operations[169] == [None] * 9
    assert {o[7] for o in operations[:169]} == {"oui", "non"}
    assert all(isinstance(o[4], Decimal) for o in operations[:169])
    assert Decimal("-2659.16") in {o[4] for o in operations[:169]}
    premiere = onglet.cell(row=entete + 1, column=1)
    assert isinstance(premiere.value, date) and premiere.number_format == FORMAT_DATE
    # Le filtre masque des lignes entières : aucune autre donnée sur les lignes des opérations
    assert all(
        onglet.cell(row=n, column=c).value is None
        for n in range(entete, derniere + 1)
        for c in range(10, 30)
    )
    titres = [g.title.tx.rich.p[0].r[0].t for g in onglet._charts]
    assert titres == [
        "Vue joint : répartition par catégorie",
        "Vue perso : répartition par catégorie",
    ]
    # Graphiques à droite des récapitulatifs, au-dessus des opérations
    assert [g.anchor for g in onglet._charts] == [f"O{LIGNE_RECAP}", f"W{LIGNE_RECAP}"]
    assert entete > LIGNE_RECAP + LIGNES_GRAPHIQUE


def test_secteurs_sur_les_lignes_de_categorie(classeur: Workbook) -> None:
    onglet = classeur["2026-07"]
    lignes = [
        n
        for n in range(LIGNE_RECAP + 2, 60)
        if onglet.cell(row=n, column=1).value not in (None, "Total")
        and not onglet.cell(row=n, column=1).font.i
        and onglet.cell(row=n, column=1).value != "Catégorie"
    ]
    lignes = lignes[:10]  # les 10 catégories du joint de juillet
    [serie] = onglet._charts[0].series
    assert serie.val.numRef.f == "(" + ",".join(f"'2026-07'!$B${n}" for n in lignes) + ")"
    assert serie.cat.strRef.f == "(" + ",".join(f"'2026-07'!$A${n}" for n in lignes) + ")"


def test_categorie_negative_exclue_des_secteurs(juillet: ContenuPublication) -> None:
    mois = juillet.mois[-1]
    vue = mois.synthese.agregats.joint
    lignes = list(vue.categories)
    lignes[1] = lignes[1].model_copy(update={"montant": Decimal("-12.00"), "sous_categories": ()})
    agregats = mois.synthese.agregats.model_copy(
        update={"joint": vue.model_copy(update={"categories": tuple(lignes)})}
    )
    synthese = mois.synthese.model_copy(update={"agregats": agregats})
    modifie = replace(juillet, mois=[*juillet.mois[:-1], replace(mois, synthese=synthese)])
    onglet = construire_classeur(modifie, MAINTENANT)["2026-07"]
    [serie] = onglet._charts[0].series
    negative = next(
        n for n in range(LIGNE_RECAP, 60) if onglet.cell(row=n, column=1).value == "Courses"
    )
    assert onglet.cell(row=negative, column=2).value == Decimal("-12.00")  # reste au récap
    assert f"$B${negative}," not in serie.val.numRef.f
    assert serie.val.numRef.f.count("$B$") == 9


def test_fichier_relu(classeur: Workbook, tmp_path: Path) -> None:
    chemin = enregistrer(classeur, tmp_path / "exports" / "budget.xlsx")
    assert not list(chemin.parent.glob(".*.tmp.xlsx"))
    relu = load_workbook(chemin)
    assert relu.sheetnames == [SYNTHESE, "2026-07", "2026-06"]
    assert relu[SYNTHESE]["A5"].value == "2026-06"
    entete, derniere = _operations(relu["2026-07"])
    assert derniere - entete == 169
    assert isinstance(relu["2026-07"].cell(row=entete + 1, column=1).value, datetime)


def test_republier_donne_le_meme_classeur(juillet: ContenuPublication, tmp_path: Path) -> None:
    chemin = tmp_path / "budget.xlsx"

    def valeurs() -> dict[str, list[tuple[object, ...]]]:
        relu = load_workbook(chemin)
        return {o.title: list(o.iter_rows(values_only=True)) for o in relu.worksheets}

    enregistrer(construire_classeur(juillet, MAINTENANT), chemin)
    premiere = valeurs()
    enregistrer(construire_classeur(juillet, MAINTENANT), chemin)
    assert valeurs() == premiere


def test_fichier_ouvert_dans_excel(
    classeur: Workbook, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def verrouille(source: object, cible: object) -> None:
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(os, "replace", verrouille)
    with pytest.raises(PublicationError, match="le fermer dans Excel"):
        enregistrer(classeur, tmp_path / "budget.xlsx")
    assert list(tmp_path.iterdir()) == []


def test_classeur_d_un_mois_sans_les_mois_suivants(
    juillet: ContenuPublication, ledger: Ledger, contexte: Contexte
) -> None:
    """Juillet est synthétisé, mais le classeur de juin s'arrête à juin."""
    juin = contenu(ledger, contexte, "2026-06")
    assert [m.mois for m in juin.mois] == ["2026-06"]
    assert construire_classeur(juin, MAINTENANT).sheetnames == [SYNTHESE, "2026-06"]


def _juin_non_classe(chemin: Path) -> None:
    with Ledger(chemin) as ledger:
        importer_mois("2026-06", DOSSIER_FIXTURES, ledger, date(2026, 10, 15))


def test_controle_bloquant_aucun_fichier(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Juin importé sans réponse : 126 opérations en attente, rien n'est écrit."""
    chemin = tmp_path / "budget.sqlite"
    _juin_non_classe(chemin)
    monkeypatch.setattr(cli.console_erreur, "width", 200)
    sortie = tmp_path / "exports"
    arguments = ["--base", str(chemin), "--contexte", str(EXEMPLE_CONTEXTE), "--sans-llm"]
    resultat = CliRunner().invoke(
        cli.app, ["publish", "--mois", "2026-06", *arguments, "--sortie", str(sortie)]
    )
    assert resultat.exit_code == 1
    assert "126 opération(s) en attente" in " ".join(resultat.output.split())
    assert not sortie.exists()


def test_mois_deja_synthetise_en_echec_bloque_la_publication(
    tmp_path: Path, contexte: Contexte
) -> None:
    """Une règle désactivée remet des opérations de juin en attente : juillet n'est pas publié."""
    with preparer_ledger(tmp_path / "budget.sqlite", contexte) as ledger:
        [regle] = [r for r in ledger.regles() if r.valeur == "XXXXXXXXXXX ECHEANCE"]
        assert regle.id is not None
        ledger.desactiver_regle(regle.id)
        with pytest.raises(ControleError, match="Contrôles bloquants du mois 2026-06"):
            contenu(ledger, contexte)


def test_commande_publish_un_fichier_par_mois(
    ledger: Ledger, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(cli.console, "width", 200)
    arguments = ["--base", str(ledger.chemin), "--contexte", str(EXEMPLE_CONTEXTE), "--sans-llm"]
    arguments += ["--sortie", str(tmp_path)]
    for mois in ("2026-06", "2026-07"):
        resultat = CliRunner().invoke(cli.app, ["publish", "--mois", mois, *arguments])
        assert resultat.exit_code == 0, resultat.output
    sortie_cli = " ".join(resultat.output.split())
    assert "2026-06 │ 7 566,46 € │ 1 322,86 € │ 137 │ oui" in sortie_cli
    assert "2026-07 │ 9 404,27 € │ 1 154,49 € │ 169 │ oui" in sortie_cli
    assert f"Classeur écrit : {tmp_path / 'budget_2026-07.xlsx'}" in sortie_cli
    assert sorted(p.name for p in tmp_path.iterdir()) == [
        "budget_2026-06.xlsx",
        "budget_2026-07.xlsx",
    ]
    assert load_workbook(tmp_path / "budget_2026-06.xlsx").sheetnames == [SYNTHESE, "2026-06"]
    assert load_workbook(tmp_path / "budget_2026-07.xlsx").sheetnames == [
        SYNTHESE,
        "2026-07",
        "2026-06",
    ]
